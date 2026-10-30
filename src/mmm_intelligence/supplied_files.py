"""No paths, network, SQL, macros, pickle or automatic semantic inference."""

import base64
import csv
import hashlib
import io
import json
import math
import zipfile

from .errors import MMMError


def unique_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise MMMError("SCHEMA_INVALID", "Duplicate JSON field names are ambiguous")
        result[key] = value
    return result


def parse_file(request, settings):
    if len(request.content_base64) > 4 * ((settings.max_file_bytes + 2) // 3):
        raise MMMError("FILE_LIMIT", "Encoded file exceeds host limit")
    try:
        raw = base64.b64decode(request.content_base64, validate=True)
    except ValueError as exc:
        raise MMMError("FILE_INVALID", "Invalid base64") from exc
    if not raw or len(raw) > settings.max_file_bytes:
        raise MMMError("FILE_LIMIT", "Empty or excessive file")
    sheets = []
    if request.format in {"xls", "parquet", "xml"}:
        from .file_readers import read_additional

        rows, sheets = read_additional(raw, request, settings)
    elif request.format == "xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                infos = archive.infolist()
                if (
                    len(infos) > 1000
                    or sum(i.file_size for i in infos) > settings.max_file_bytes * 4
                ):
                    raise MMMError("FILE_LIMIT", "Excessive expanded workbook")
                if any(
                    ".." in i.filename.split("/")
                    or i.filename.startswith(("/", "\\"))
                    or "vba" in i.filename.lower()
                    or "externallinks/" in i.filename.lower()
                    for i in infos
                ):
                    raise MMMError("ACTIVE_CONTENT", "Unsafe workbook content")
            try:
                import openpyxl

                if not openpyxl.DEFUSEDXML:
                    raise MMMError("FORMAT_UNAVAILABLE", "XLSX requires defusedxml protection")
            except ImportError as exc:
                raise MMMError(
                    "FORMAT_UNAVAILABLE", "Install the optional files dependencies for xlsx"
                ) from exc
            book = openpyxl.load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheets = book.sheetnames
                if request.sheet is None and len(sheets) != 1:
                    raise MMMError("SHEET_REQUIRED", "Select an explicit worksheet")
                if request.sheet is not None and request.sheet not in sheets:
                    raise MMMError("SHEET_INVALID", "Unknown worksheet")
                sheet = book[request.sheet or sheets[0]]
                if (sheet.max_row or 0) > settings.max_rows + request.header_row or (
                    sheet.max_column or 0
                ) > 80:
                    raise MMMError("ROW_LIMIT", "Workbook dimensions exceed host limits")
                table = []
                sheet.reset_dimensions()
                for row in sheet.iter_rows():
                    if len(table) >= settings.max_rows + request.header_row or len(row) > 80:
                        raise MMMError("ROW_LIMIT", "Actual workbook size exceeds limits")
                    if any(c.data_type == "f" for c in row):
                        raise MMMError("ACTIVE_CONTENT", "Spreadsheet formulas are not accepted")
                    from datetime import date, datetime

                    table.append(
                        [
                            (
                                c.value.date().isoformat()
                                if isinstance(c.value, datetime)
                                and c.value.time() == datetime.min.time()
                                else c.value.isoformat()
                            )
                            if isinstance(c.value, (date, datetime))
                            else c.value
                            for c in row
                        ]
                    )
            finally:
                book.close()
        except (zipfile.BadZipFile, KeyError) as exc:
            raise MMMError("FILE_INVALID", "Malformed workbook") from exc
        rows = tabular(table, request.header_row)
    else:
        text = raw.decode(request.encoding)
        if request.format in {"csv", "tsv"}:
            table = []
            for row in csv.reader(
                io.StringIO(text),
                delimiter=request.delimiter or ("\t" if request.format == "tsv" else ","),
                strict=True,
            ):
                if len(table) > settings.max_rows + request.header_row:
                    raise MMMError("ROW_LIMIT", "Too many source rows")
                table.append(row)
            rows = tabular(table, request.header_row)
        elif request.format == "json":
            rows = json.loads(text, object_pairs_hook=unique_object)
        else:
            rows = [
                json.loads(line, object_pairs_hook=unique_object)
                for line in text.splitlines()
                if line.strip()
            ]
    if not isinstance(rows, list) or not rows or len(rows) > settings.max_rows:
        raise MMMError("ROW_LIMIT", "Expected a nonempty bounded array of records")
    columns = list(rows[0]) if isinstance(rows[0], dict) else []
    if (
        not columns
        or len(columns) > 80
        or any(not isinstance(c, str) or not c or len(c) > 128 for c in columns)
    ):
        raise MMMError("SCHEMA_INVALID", "Invalid columns")
    for row in rows:
        if not isinstance(row, dict) or set(row) != set(columns):
            raise MMMError(
                "SCHEMA_DRIFT", "All rows must declare the same fields; use null for unknown"
            )
        for value in row.values():
            if (
                not isinstance(value, (str, int, float, bool, type(None)))
                or isinstance(value, float)
                and not math.isfinite(value)
            ):
                raise MMMError("SCHEMA_INVALID", "Only finite scalar cells are supported")
            if isinstance(value, str) and len(value) > 10000:
                raise MMMError("CELL_LIMIT", "Cell text exceeds limit")
    return raw, rows, columns, sheets, hashlib.sha256(raw).hexdigest()


def tabular(table, header_row):
    if len(table) < header_row:
        raise MMMError("HEADER_INVALID", "Missing header")
    header = table[header_row - 1]
    if any(not isinstance(c, str) or not c for c in header) or len(set(header)) != len(header):
        raise MMMError("HEADER_INVALID", "Headers must be unique nonempty strings")
    rows = []
    for cells in table[header_row:]:
        if len(cells) != len(header):
            raise MMMError("ROW_INVALID", "Cell count differs from header")
        rows.append(dict(zip(header, [None if c == "" else c for c in cells])))
    return rows
